"""Build the Excel workbook from the public exports (Week 12). Called by scripts/export_bi.py.

The workbook holds the public tables as Excel Tables and computes everything else with
formulas (XLOOKUP, COUNTIFS), so it recalculates when a selection changes. Sheets:

  README          sources, the publication rule, how to use the sheets
  Data            rpt_public_company_year (Excel Table tblCompanyYear)
  Peer Benchmark  pick a company, a KPI and a year: value, peer median, rank, percentile
  Sector Summary  group figures from rpt_public_group_metrics; withheld cells show "n<5"
  DQS view        Disclosure Quality Score of every company in a chosen year
  DQS Data, Group Data, Lists   the tables the formulas read (and the ranking rules)

Statistics are never recomputed from company rows: medians and group figures come from
rpt_public_group_metrics, where the 5-company rule is already applied. Rank and percentile are
positions within the ranking group, computed with the warehouse's own method (percent rank,
ties share the better place, only for 5 or more companies and a material KPI).
"""

from __future__ import annotations

import csv
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parent.parent
SEEDS = ROOT / "dbt" / "seeds"

FONT = "Arial"
BOLD = Font(name=FONT, bold=True)
TITLE = Font(name=FONT, bold=True, size=14)
NORMAL = Font(name=FONT)
GREY = Font(name=FONT, italic=True, color="666666")
INPUT_FILL = PatternFill("solid", start_color="FFF2CC")  # pale yellow: cells to change
HEADER_FILL = PatternFill("solid", start_color="DDEBF7")
WITHHELD_FILL = PatternFill("solid", start_color="EDEDED")
WRAP = Alignment(wrap_text=True, vertical="top")

# KPIs of the benchmark tool: those with a published peer median in rpt_public_group_metrics.
BENCHMARK_KPIS = [
    # label, company column, median column, count column, direction, unit
    (
        "GHG intensity (Scope 1+2 per Rs crore)",
        "ghg_intensity_tco2e_per_cr",
        "ghg_intensity_median",
        "companies_with_ghg_intensity",
        "lower_better",
        "tCO2e per Rs crore",
    ),
    (
        "Energy intensity (per Rs crore)",
        "energy_intensity_gj_per_cr",
        "energy_intensity_median",
        "companies_with_energy_intensity",
        "lower_better",
        "GJ per Rs crore",
    ),
    (
        "Water intensity (per Rs crore)",
        "water_intensity_kl_per_cr",
        "water_intensity_median",
        "companies_with_water_intensity",
        "lower_better",
        "kl per Rs crore",
    ),
    (
        "Waste intensity (per Rs crore)",
        "waste_intensity_t_per_cr",
        "waste_intensity_median",
        "companies_with_waste_intensity",
        "lower_better",
        "t per Rs crore",
    ),
]

# Sector Summary columns: (header, column in rpt_public_group_metrics, number format, is a
# restatement figure that can be "not yet observable", is a plain count)
SUMMARY_COLUMNS = [
    ("Companies reporting", "companies_reporting", "0", False, True),
    ("GHG intensity, whole group", "ghg_intensity_agg", "#,##0.00", False, False),
    ("GHG intensity, typical company", "ghg_intensity_median", "#,##0.00", False, False),
    ("Energy intensity, whole group", "energy_intensity_agg", "#,##0.0", False, False),
    ("Energy intensity, typical company", "energy_intensity_median", "#,##0.0", False, False),
    ("Water intensity, whole group", "water_intensity_agg", "#,##0.0", False, False),
    ("Water intensity, typical company", "water_intensity_median", "#,##0.0", False, False),
    ("Waste intensity, whole group", "waste_intensity_agg", "#,##0.00", False, False),
    ("Waste intensity, typical company", "waste_intensity_median", "#,##0.00", False, False),
    ("Renewable share, whole group (%)", "renewable_share_agg", "0.0", False, False),
    ("Women in workforce, whole group (%)", "female_workforce_agg", "0.0", False, False),
    ("Scope 3 disclosed (% of reports)", "scope3_disclosed_share", "0.0", False, False),
    ("DQS, median", "dqs_median", "0.0", False, False),
    ("Grade A (% of graded reports)", "dqs_grade_a_share", "0.0", False, False),
    ("Figures restated materially (%)", "material_restatement_share", "0.0", True, False),
]
SUMMARY_GROUP_TYPES = ["all", "sector_group", "peer_group", "ownership_type"]

DQS_COLUMNS = [
    ("DQS (0-100)", "dqs_total", "0.0"),
    ("Grade", "dqs_grade", "@"),
    ("Rank in group", "dqs_rank_in_group", "0"),
    ("Percentile in group", "dqs_percentile_in_group", "0.0"),
    ("Provisional", "is_provisional", "@"),
    ("Assurance level", "assurance_level", "@"),
    ("Assurance scope", "assurance_scope", "@"),
    ("Reason codes", "dqs_reason_codes", "@"),
]


# ------------------------------------------------------------------ helpers


def _clean(value):
    if value is None:
        return None
    if isinstance(value, float) and pd.isna(value):
        return None
    if value is pd.NA or value is pd.NaT:
        return None
    if hasattr(value, "item"):  # numpy scalar
        return value.item()
    return value


def _write_table(
    ws, frame: pd.DataFrame, name: str, key_formula=None, first_row: int = 1, freeze: bool = True
):
    """Write `frame` as an Excel Table. `key_formula(row_number)` adds a lookup_key column."""
    columns = list(frame.columns) + (["lookup_key"] if key_formula else [])
    for c, column in enumerate(columns, start=1):
        cell = ws.cell(row=first_row, column=c, value=column)
        cell.font = BOLD
    for r, record in enumerate(frame.itertuples(index=False), start=first_row + 1):
        for c, value in enumerate(record, start=1):
            ws.cell(row=r, column=c, value=_clean(value)).font = NORMAL
        if key_formula:
            ws.cell(row=r, column=len(columns), value=key_formula(r)).font = NORMAL
    last_row = first_row + max(len(frame), 1)
    ref = f"A{first_row}:{get_column_letter(len(columns))}{last_row}"
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(table)
    if freeze:
        ws.freeze_panes = ws.cell(row=first_row + 1, column=1)
    for c, column in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(c)].width = min(max(len(column) + 2, 12), 40)
    return {column: get_column_letter(c) for c, column in enumerate(columns, start=1)}


def _name(wb: Workbook, name: str, ref: str) -> None:
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


def _list_validation(ws, cell: str, source: str) -> None:
    validation = DataValidation(type="list", formula1=source, allow_blank=False)
    validation.error = "Pick a value from the list."
    validation.showErrorMessage = True
    ws.add_data_validation(validation)
    validation.add(ws[cell])


def _input(ws, cell: str, value) -> None:
    ws[cell] = value
    ws[cell].fill = INPUT_FILL
    ws[cell].font = BOLD


def _label(ws, cell: str, text: str, font=NORMAL) -> None:
    ws[cell] = text
    ws[cell].font = font


def ranking_rules(company_year: pd.DataFrame) -> pd.DataFrame:
    """Is each benchmark KPI ranked in each ranking group? Material for every peer group of the
    ranking group (dbt/seeds/metric_materiality.csv, through kpi_catalogue.materiality_metric_id),
    the rule of fct_company_kpi_percentile. These are rules, not data."""
    with (SEEDS / "kpi_catalogue.csv").open(encoding="utf-8") as fh:
        catalogue = {row["kpi_name"]: row["materiality_metric_id"] for row in csv.DictReader(fh)}
    with (SEEDS / "metric_materiality.csv").open(encoding="utf-8") as fh:
        material = {
            (row["peer_group"], row["metric_id"]): row["is_material"].strip().lower() == "true"
            for row in csv.DictReader(fh)
        }
    groups = company_year[["ranking_group", "peer_group"]].dropna().drop_duplicates()
    rows = []
    for ranking_group, peers in groups.groupby("ranking_group"):
        for _, column, *_ in BENCHMARK_KPIS:
            metric = catalogue.get(column, "")
            ranked = all(material.get((peer, metric), False) for peer in peers["peer_group"])
            rows.append(
                {
                    "key": f"{ranking_group}|{column}",
                    "ranking_group": ranking_group,
                    "kpi": column,
                    "material": "yes" if ranked else "no",
                }
            )
    return pd.DataFrame(rows, columns=["key", "ranking_group", "kpi", "material"])


# ------------------------------------------------------------------ sheets


def _readme(ws, built: str, local: bool) -> None:
    ws.column_dimensions["A"].width = 110
    lines = [
        (
            "ESG Lens India - " + ("LOCAL workbook (never share)" if local else "public workbook"),
            TITLE,
        ),
        (
            f"Built {built} from the warehouse of ESG Lens India (NIFTY 50 BRSR filings, FY2022-23 to FY2025-26).",
            NORMAL,
        ),
        ("", NORMAL),
        ("Sources", BOLD),
        (
            "Data, Group Data and DQS Data are copies of the public models rpt_public_company_year, rpt_public_group_metrics and rpt_public_disclosure_quality.",
            NORMAL,
        ),
        (
            'Every figure goes back to a company\'s BRSR filing. Cite it in plain text, for example: "HDFC Bank Limited, BRSR FY2025-26, filed on NSE (date)". No links to the exchange.',
            NORMAL,
        ),
        (
            "Method and definitions: docs/kpi_definitions.md, docs/semantic_layer.md and docs/disclosure_quality_score.md in the project repository.",
            NORMAL,
        ),
        ("", NORMAL),
        ("Publication rule", BOLD),
        (
            "Only values ESG Lens India calculates are in this workbook: intensities, shares, gaps, medians, ranks, percentiles, scores and counts. Filed values (tonnes, GJ, rupees, headcounts, filed percentages) are not.",
            NORMAL,
        ),
        (
            'A group statistic (median, whole-group ratio, share) is withheld when fewer than 5 companies stand behind it. Withheld cells show "n<5". Do not fill them in from company rows.',
            NORMAL,
        ),
        (
            '"not yet observable" = a restatement figure of the latest year: the next report, which could restate it, does not exist yet.',
            NORMAL,
        ),
        (
            "FY2025-26 is provisional: its figures and scores can still change when the FY2026-27 reports arrive.",
            NORMAL,
        ),
        (
            "The Disclosure Quality Score is not an ESG rating, and this workbook is not investment advice.",
            BOLD,
        ),
        ("", NORMAL),
        ("How to use", BOLD),
        (
            "Cells with a pale yellow fill are the only ones to change. Pick from the drop-down lists; everything else is a formula and recalculates.",
            NORMAL,
        ),
        (
            "Peer Benchmark: choose a company, a KPI and a year. Shows the company's value, the median of its ranking group, its rank (1 = best) and its percentile (100 = best).",
            NORMAL,
        ),
        (
            "Sector Summary: choose a year. Whole group = all emissions over all turnover of the group (large companies weigh more); typical company = the median.",
            NORMAL,
        ),
        (
            "DQS view: choose a year. Compare scores within one year only: the years are built from different parts.",
            NORMAL,
        ),
        (
            "Rank and percentile follow the warehouse rule: only when 5 or more companies of the ranking group have a value and the KPI is material for the group (Lists, tblRules).",
            NORMAL,
        ),
    ]
    if local:
        lines.insert(
            2,
            (
                "LOCAL ONLY: the sheet Local Data holds raw filed values (fct_company_year). Do not share, upload or screenshot this file.",
                BOLD,
            ),
        )
    for r, (text, font) in enumerate(lines, start=1):
        cell = ws.cell(row=r, column=1, value=text)
        cell.font = font
        cell.alignment = WRAP


def _lists(wb, ws, company_year: pd.DataFrame, rules: pd.DataFrame) -> dict[str, dict[str, str]]:
    kpis = pd.DataFrame(
        BENCHMARK_KPIS,
        columns=["label", "column", "median_column", "count_column", "direction", "unit"],
    )
    _write_table(ws, kpis, "tblKpi", freeze=False)
    start = len(kpis) + 4
    ws.cell(
        row=start - 1,
        column=1,
        value="Ranking rules (material for every peer group of the ranking group)",
    ).font = BOLD
    _write_table(ws, rules, "tblRules", first_row=start, freeze=False)

    companies = sorted(company_year["display_name"].dropna().unique())
    years = sorted(company_year["fiscal_year_label"].dropna().unique())
    for column, header, values in (("H", "Companies", companies), ("I", "Years", years)):
        ws[f"{column}1"] = header
        ws[f"{column}1"].font = BOLD
        for r, value in enumerate(values, start=2):
            ws[f"{column}{r}"] = value
    _name(wb, "CompanyList", f"Lists!$H$2:$H${len(companies) + 1}")
    _name(wb, "YearList", f"Lists!$I$2:$I${len(years) + 1}")
    _name(wb, "KpiList", f"Lists!$A$2:$A${len(kpis) + 1}")
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["H"].width = 30
    return {"companies": companies, "years": years}


def _peer_benchmark(ws, default_company: str, latest_year: str) -> None:
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 70
    _label(ws, "A1", "Peer Benchmark", TITLE)
    _label(ws, "A2", "Change the yellow cells. Everything else is a formula.", GREY)
    _label(ws, "A4", "Company", BOLD)
    _label(ws, "A5", "KPI", BOLD)
    _label(ws, "A6", "Financial year", BOLD)
    _input(ws, "B4", default_company)
    _input(ws, "B5", BENCHMARK_KPIS[0][0])
    _input(ws, "B6", latest_year)
    _list_validation(ws, "B4", "CompanyList")
    _list_validation(ws, "B5", "KpiList")
    _list_validation(ws, "B6", "YearList")

    # helper cells (column E, hidden): which column of tblCompanyYear, which direction, ...
    helpers = {
        "E4": '=$B$4&"|"&$B$6',
        "E5": "=XLOOKUP($B$5,tblKpi[label],tblKpi[column])",
        "E6": "=MATCH(E5,tblCompanyYear[#Headers],0)",
        "E7": "=XLOOKUP($B$5,tblKpi[label],tblKpi[direction])",
        "E8": "=MATCH(XLOOKUP($B$5,tblKpi[label],tblKpi[median_column]),tblGroup[#Headers],0)",
        "E9": "=MATCH(XLOOKUP($B$5,tblKpi[label],tblKpi[count_column]),tblGroup[#Headers],0)",
        "E10": '="ranking_group|"&B9&"|"&$B$6',
        "E11": '=XLOOKUP(B9&"|"&E5,tblRules[key],tblRules[material],"no")',
    }
    for cell, formula in helpers.items():
        ws[cell] = formula
    ws.column_dimensions["E"].hidden = True

    rows = [
        (
            "A9",
            "Ranking group",
            "B9",
            '=XLOOKUP(E4,tblCompanyYear[lookup_key],tblCompanyYear[ranking_group],"no report this year")',
            "@",
            "The group the company is compared with (its peer group, or its sector group when the peer group is small).",
        ),
        (
            "A10",
            "Company value",
            "B10",
            '=IF(B9="no report this year","-",IF(ISBLANK(XLOOKUP(E4,tblCompanyYear[lookup_key],INDEX(tblCompanyYear[#Data],0,E6))),"no value",XLOOKUP(E4,tblCompanyYear[lookup_key],INDEX(tblCompanyYear[#Data],0,E6))))',
            "#,##0.00",
            "Computed by ESG Lens India from the filing (an intensity per Rs crore of turnover).",
        ),
        ("A11", "Unit", "B11", "=XLOOKUP($B$5,tblKpi[label],tblKpi[unit])", "@", ""),
        (
            "A12",
            "Companies in the group with a value",
            "B12",
            '=IF(B9="no report this year",0,COUNTIFS(tblCompanyYear[ranking_group],B9,tblCompanyYear[fiscal_year_label],$B$6,INDEX(tblCompanyYear[#Data],0,E6),"<>"))',
            "0",
            "A rank and a percentile need 5 or more.",
        ),
        (
            "A13",
            "Ranked in this group?",
            "B13",
            '=IF(NOT(ISNUMBER(B10)),"no value",IF(E11<>"yes","not ranked: not material for this group",IF(B12<5,"not ranked: n<5","yes")))',
            "@",
            "Material = the KPI matters for every peer group of the ranking group (Lists, tblRules).",
        ),
        (
            "A14",
            "Rank (1 = best)",
            "B14",
            '=IF(B13<>"yes","-",1+IF(E7="lower_better",COUNTIFS(tblCompanyYear[ranking_group],B9,tblCompanyYear[fiscal_year_label],$B$6,INDEX(tblCompanyYear[#Data],0,E6),"<"&B10),COUNTIFS(tblCompanyYear[ranking_group],B9,tblCompanyYear[fiscal_year_label],$B$6,INDEX(tblCompanyYear[#Data],0,E6),">"&B10)))',
            "0",
            "Lower intensity is better. Companies with the same value share the better rank.",
        ),
        (
            "A15",
            "Percentile in group (100 = best)",
            "B15",
            '=IF(B13<>"yes","-",ROUND(100*IF(E7="lower_better",COUNTIFS(tblCompanyYear[ranking_group],B9,tblCompanyYear[fiscal_year_label],$B$6,INDEX(tblCompanyYear[#Data],0,E6),">"&B10),COUNTIFS(tblCompanyYear[ranking_group],B9,tblCompanyYear[fiscal_year_label],$B$6,INDEX(tblCompanyYear[#Data],0,E6),"<"&B10))/(B12-1),1))',
            "0.0",
            "Share of the other companies in the group that do worse (percent rank, as in the warehouse).",
        ),
        (
            "A16",
            "Peer median (typical company of the group)",
            "B16",
            '=IF(COUNTIF(tblGroup[lookup_key],E10)=0,"-",IF(ISBLANK(XLOOKUP(E10,tblGroup[lookup_key],INDEX(tblGroup[#Data],0,E8))),"n<5",XLOOKUP(E10,tblGroup[lookup_key],INDEX(tblGroup[#Data],0,E8))))',
            "#,##0.00",
            "From rpt_public_group_metrics: withheld (n<5) when fewer than 5 companies stand behind it.",
        ),
        (
            "A17",
            "Companies behind the median",
            "B17",
            '=IF(COUNTIF(tblGroup[lookup_key],E10)=0,"-",XLOOKUP(E10,tblGroup[lookup_key],INDEX(tblGroup[#Data],0,E9)))',
            "0",
            "",
        ),
        (
            "A18",
            "Company against peer median",
            "B18",
            '=IF(AND(ISNUMBER(B10),ISNUMBER(B16)),B10/B16-1,"-")',
            "+0%;-0%;0%",
            "Above 0% = more intensive than the typical company of its group.",
        ),
    ]
    for label_cell, label, value_cell, formula, number_format, note in rows:
        _label(ws, label_cell, label, BOLD)
        ws[value_cell] = formula
        ws[value_cell].number_format = number_format
        ws[value_cell].font = NORMAL
        ws[value_cell].alignment = Alignment(horizontal="right")
        if note:
            _label(ws, "C" + value_cell[1:], note, GREY)
    provisional = '=IF($B$6=INDEX(YearList,ROWS(YearList)),"Provisional year: no later report has confirmed these figures yet.","")'
    ws["A20"] = provisional
    ws["A20"].font = GREY
    _label(
        ws,
        "A22",
        "The Disclosure Quality Score is not an ESG rating, and this workbook is not investment advice.",
        GREY,
    )


def _sector_summary(ws, group_values: dict[str, list[str]], latest_year: str) -> None:
    _label(ws, "A1", "Sector Summary", TITLE)
    _label(
        ws,
        "A2",
        'Group figures from rpt_public_group_metrics. "n<5" = withheld, fewer than 5 companies; "not yet observable" = no later report yet.',
        GREY,
    )
    _label(ws, "A3", "Financial year", BOLD)
    _input(ws, "B3", latest_year)
    _list_validation(ws, "B3", "YearList")
    header_row = 5
    _label(ws, f"A{header_row}", "Group type", BOLD)
    _label(ws, f"B{header_row}", "Group", BOLD)
    for c, (header, *_rest) in enumerate(SUMMARY_COLUMNS, start=3):
        cell = ws.cell(row=header_row, column=c, value=header)
        cell.font = BOLD
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[get_column_letter(c)].width = 14
    ws.row_dimensions[header_row].height = 45
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 26
    r = header_row + 1
    for group_type in SUMMARY_GROUP_TYPES:
        for group in group_values.get(group_type, []):
            ws.cell(row=r, column=1, value=group_type).font = GREY
            ws.cell(row=r, column=2, value=group).font = NORMAL
            key = f'$A{r}&"|"&$B{r}&"|"&$B$3'
            for c, (_header, column, number_format, restatement, count) in enumerate(
                SUMMARY_COLUMNS, start=3
            ):
                lookup = f"XLOOKUP({key},tblGroup[lookup_key],tblGroup[{column}])"
                if count:
                    formula = f'=IF(COUNTIF(tblGroup[lookup_key],{key})=0,"-",{lookup})'
                else:
                    reason = '"n<5"'
                    if restatement:
                        reason = (
                            f'IF(ISNUMBER(SEARCH("not_yet",XLOOKUP({key},tblGroup[lookup_key],'
                            f'tblGroup[suppressed_reason]))),"not yet observable","n<5")'
                        )
                    formula = (
                        f'=IF(COUNTIF(tblGroup[lookup_key],{key})=0,"-",'
                        f"IF(ISBLANK({lookup}),{reason},{lookup}))"
                    )
                cell = ws.cell(row=r, column=c, value=formula)
                cell.number_format = number_format
                cell.font = NORMAL
            r += 1
    last = r - 1
    area = f"C{header_row + 1}:{get_column_letter(2 + len(SUMMARY_COLUMNS))}{last}"
    ws.conditional_formatting.add(
        area,
        FormulaRule(formula=[f"ISTEXT(C{header_row + 1})"], fill=WITHHELD_FILL, font=GREY),
    )
    ws.freeze_panes = f"C{header_row + 1}"
    _label(
        ws,
        f"A{last + 2}",
        "Whole group = sum of the numerators over sum of the denominators of the companies that have both (large companies weigh more). Typical company = the median. Never an average of company ratios.",
        GREY,
    )
    _label(
        ws,
        f"A{last + 3}",
        "The Disclosure Quality Score is not an ESG rating, and this workbook is not investment advice.",
        GREY,
    )


def _dqs_view(ws, companies: list[str], latest_final_year: str) -> None:
    _label(ws, "A1", "DQS view", TITLE)
    _label(
        ws,
        "A2",
        "Disclosure Quality Score of every company in one year. Compare within a year only.",
        GREY,
    )
    _label(ws, "A3", "Financial year", BOLD)
    _input(ws, "B3", latest_final_year)
    _list_validation(ws, "B3", "YearList")
    ws["C3"] = (
        "=IF(COUNTIFS(tblDqs[fiscal_year_label],$B$3,tblDqs[is_provisional],TRUE)>0,"
        '"PROVISIONAL: no later report yet, scores can still fall","Final scores")'
    )
    ws["C3"].font = BOLD
    _label(ws, "A4", "Median score, all companies", BOLD)
    ws["B4"] = (
        '=IF(ISBLANK(XLOOKUP("all|All companies|"&$B$3,tblGroup[lookup_key],tblGroup[dqs_median])),'
        '"n<5",XLOOKUP("all|All companies|"&$B$3,tblGroup[lookup_key],tblGroup[dqs_median]))'
    )
    ws["B4"].number_format = "0.0"
    header_row = 6
    headers = ["Company"] + [header for header, *_ in DQS_COLUMNS]
    for c, header in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=c, value=header)
        cell.font = BOLD
        cell.fill = HEADER_FILL
    for r, company in enumerate(companies, start=header_row + 1):
        ws.cell(row=r, column=1, value=company).font = NORMAL
        key = f'$A{r}&"|"&$B$3'
        for c, (_header, column, number_format) in enumerate(DQS_COLUMNS, start=2):
            lookup = f"XLOOKUP({key},tblDqs[lookup_key],tblDqs[{column}])"
            formula = (
                f'=IF(COUNTIF(tblDqs[lookup_key],{key})=0,"no report",'
                f'IF(ISBLANK({lookup}),"-",{lookup}))'
            )
            cell = ws.cell(row=r, column=c, value=formula)
            cell.number_format = number_format
            cell.font = NORMAL
    ws.column_dimensions["A"].width = 30
    for c in range(2, len(headers) + 1):
        ws.column_dimensions[get_column_letter(c)].width = 14
    ws.column_dimensions[get_column_letter(len(headers))].width = 60
    ws.freeze_panes = f"B{header_row + 1}"
    last = header_row + len(companies)
    _label(
        ws,
        f"A{last + 2}",
        "Rank and percentile are within the ranking group and year, only with 5 or more gradable companies. The Disclosure Quality Score is not an ESG rating, and this workbook is not investment advice.",
        GREY,
    )


# ------------------------------------------------------------------ the workbook


def _prefix_new_functions(ws) -> None:
    """Excel stores functions newer than Excel 2007 with an _xlfn. prefix (its interface hides
    it). openpyxl writes formulas as given, so add the prefix, or Excel shows #NAME?."""
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.value = cell.value.replace("_xlfn.XLOOKUP(", "XLOOKUP(").replace(
                    "XLOOKUP(", "_xlfn.XLOOKUP("
                )


def build(public_folder: Path | None, out: Path, local_folder: Path | None = None) -> Path:
    """Build the workbook from the public Parquet files (and, for the local version, add the
    full company-year table from the local export)."""
    if public_folder is None:
        raise SystemExit("the workbook needs the public export (run without --only local)")
    company_year = pd.read_parquet(public_folder / "rpt_public_company_year.parquet")
    dqs = pd.read_parquet(public_folder / "rpt_public_disclosure_quality.parquet")
    groups = pd.read_parquet(public_folder / "rpt_public_group_metrics.parquet")

    wb = Workbook()
    readme = wb.active
    readme.title = "README"
    data = wb.create_sheet("Data")
    peer = wb.create_sheet("Peer Benchmark")
    summary = wb.create_sheet("Sector Summary")
    dqs_view = wb.create_sheet("DQS view")
    dqs_data = wb.create_sheet("DQS Data")
    group_data = wb.create_sheet("Group Data")
    lists = wb.create_sheet("Lists")

    built = date.today().isoformat()
    _readme(readme, built, local=local_folder is not None)

    company_year = company_year.sort_values(["display_name", "fiscal_year_label"])
    columns = _write_table(data, company_year, "tblCompanyYear", lambda r: None)
    name_col, year_col = columns["display_name"], columns["fiscal_year_label"]
    key_col = columns["lookup_key"]
    for r in range(2, len(company_year) + 2):
        data[f"{key_col}{r}"] = f'={name_col}{r}&"|"&{year_col}{r}'

    dqs = dqs.sort_values(["display_name", "fiscal_year_label"])
    columns = _write_table(dqs_data, dqs, "tblDqs", lambda r: None)
    for r in range(2, len(dqs) + 2):
        dqs_data[f"{columns['lookup_key']}{r}"] = (
            f'={columns["display_name"]}{r}&"|"&{columns["fiscal_year_label"]}{r}'
        )

    groups = groups.sort_values(["group_type", "group_value", "fiscal_year_label"])
    columns = _write_table(group_data, groups, "tblGroup", lambda r: None)
    for r in range(2, len(groups) + 2):
        group_data[f"{columns['lookup_key']}{r}"] = (
            f'={columns["group_type"]}{r}&"|"&{columns["group_value"]}{r}'
            f'&"|"&{columns["fiscal_year_label"]}{r}'
        )

    lists_info = _lists(wb, lists, company_year, ranking_rules(company_year))
    years = lists_info["years"]
    latest_year = years[-1] if years else ""
    final_years = sorted(dqs.loc[~dqs["is_provisional"].astype(bool), "fiscal_year_label"].unique())
    latest_final_year = final_years[-1] if final_years else latest_year
    default_company = lists_info["companies"][0] if lists_info["companies"] else ""

    _peer_benchmark(peer, default_company, latest_year)
    group_values = {
        group_type: sorted(groups.loc[groups["group_type"] == group_type, "group_value"].unique())
        for group_type in SUMMARY_GROUP_TYPES
    }
    _sector_summary(summary, group_values, latest_year)
    _dqs_view(dqs_view, lists_info["companies"], latest_final_year)

    if local_folder is not None:
        local = wb.create_sheet("Local Data", 2)
        _label(local, "A1", "LOCAL ONLY - raw filed values (fct_company_year). Never share.", BOLD)
        full = pd.read_parquet(local_folder / "fct_company_year.parquet")
        _write_table(local, full, "tblLocalCompanyYear", first_row=3)

    for ws in wb.worksheets:
        _prefix_new_functions(ws)
        ws.sheet_view.showGridLines = ws.title not in ("README", "Peer Benchmark")
    wb.active = 0
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out
