"""The BI engine (Week 12): exports keep the public folder public, the Power BI models follow
the semantic layer's rules and parse with Power BI's own library, the Excel workbook opens and
calculates without errors and agrees with the warehouse, and the DAX formulas reconcile with
mf query."""

from __future__ import annotations

import csv
import importlib.util
import json
import re
import shutil
import subprocess
import sys
from pathlib import Path

import duckdb
import pandas as pd
import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
SCRIPTS = ROOT / "scripts"
POWERBI = ROOT / "powerbi"
EXPORTS = ROOT / "exports"
WAREHOUSE = ROOT / "data" / "warehouse" / "esg_lens.duckdb"
RECONCILIATION_DOC = ROOT / "docs" / "powerbi_reconciliation.md"
GUIDE = ROOT / "docs" / "powerbi_build_guide.md"


def _load(name: str):
    if str(SCRIPTS) not in sys.path:
        sys.path.insert(0, str(SCRIPTS))
    spec = importlib.util.spec_from_file_location(name, SCRIPTS / f"{name}.py")
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


EXPORT = _load("export_bi")
EXCEL = _load("build_excel")
PBI = _load("build_powerbi")
RECONCILE = _load("reconcile_bi")


def _real_warehouse() -> bool:
    if not WAREHOUSE.exists():
        return False
    con = duckdb.connect(str(WAREHOUSE), read_only=True)
    try:
        tables = {
            row[0] for row in con.sql("select table_name from information_schema.tables").fetchall()
        }
        if not {"rpt_public_group_metrics", "fct_company_year", "dim_fiscal_year"} <= tables:
            return False
        return con.sql("select count(distinct isin) from fct_company_year").fetchone()[0] >= 40
    finally:
        con.close()


needs_warehouse = pytest.mark.skipif(not _real_warehouse(), reason="needs the real warehouse")


# ------------------------------------------------------------------ exports


def test_public_folder_refuses_anything_but_the_public_models(tmp_path):
    for name in EXPORT.PUBLIC_FILES:
        (tmp_path / name).write_text("x")
    EXPORT.assert_public_only(tmp_path)
    (tmp_path / "fct_company_year.parquet").write_text("x")
    with pytest.raises(SystemExit, match="fct_company_year"):
        EXPORT.assert_public_only(tmp_path)


def test_public_export_writes_only_rpt_public_files(tmp_path):
    con = duckdb.connect(":memory:")
    for name in EXPORT.PUBLIC_TABLES:
        con.execute(f"create table {name} as select 'x' as isin, 1.5::decimal(4,2) as value")
    (tmp_path / "public").mkdir()
    (tmp_path / "public" / "stale_local_table.parquet").write_text("left over")
    EXPORT.export_public(con, tmp_path / "public")
    written = {path.name for path in (tmp_path / "public").iterdir()}
    assert written == {f"{name}.parquet" for name in EXPORT.PUBLIC_TABLES}
    # decimals become plain doubles for Power BI and Excel
    kind = duckdb.sql(
        f"describe select * from '{(tmp_path / 'public' / 'rpt_public_group_metrics.parquet').as_posix()}'"
    ).fetchall()
    assert dict((row[0], row[1]) for row in kind)["value"] == "DOUBLE"


def test_public_export_rounds_fractions_so_rebuilds_write_the_same_bytes(tmp_path):
    """Floating-point noise (a different summation order) must not change the committed files:
    fractional columns are rounded to PUBLIC_DECIMALS, whole numbers and text are untouched,
    and the reconciliation compares with the same rounding."""
    assert EXPORT.PUBLIC_DECIMALS == RECONCILE.PUBLIC_DECIMALS == 6
    written = []
    for noise in (0.0, 1e-13):
        con = duckdb.connect(":memory:")
        for name in EXPORT.PUBLIC_TABLES:
            con.execute(
                f"create table {name} as select * from (values "
                f"('b', 2::bigint, 0.1 + 0.2 + {noise}, 85.46070349::double, 1.25::decimal(4,2)), "
                f"('a', 1::bigint, 1/3 + {noise}, null::double, 0.5::decimal(4,2))) "
                "as t(isin, n_companies, share, pct, value)"
            )
        folder = tmp_path / f"public_{noise}"
        folder.mkdir()
        EXPORT.export_public(con, folder)
        written.append({p.name: p.read_bytes() for p in folder.iterdir()})
        con.close()
    assert written[0] == written[1]
    frame = pd.read_parquet(tmp_path / "public_0.0" / "rpt_public_group_metrics.parquet")
    assert list(frame["isin"]) == ["a", "b"] and list(frame["n_companies"]) == [1, 2]
    assert list(frame["share"]) == [0.333333, 0.3] and frame["pct"].iloc[1] == 85.460703
    assert frame["pct"].isna().iloc[0] and list(frame["value"]) == [0.5, 1.25]


def test_committed_public_files_hold_at_most_six_decimals():
    public = EXPORTS / "public"
    if not (public / "rpt_public_group_metrics.parquet").exists():
        pytest.skip("run dev.ps1 bi-export first")
    for path in sorted(public.glob("rpt_public_*.parquet")):
        frame = pd.read_parquet(path)
        for column in frame.select_dtypes("float").columns:
            values = frame[column].dropna()
            drift = (values - values.round(EXPORT.PUBLIC_DECIMALS)).abs()
            assert values.empty or drift.max() < 1e-12, (path.name, column)


def test_public_tables_are_exactly_the_rpt_public_models():
    assert set(EXPORT.PUBLIC_TABLES) == {
        path.stem for path in (ROOT / "dbt" / "models" / "marts").glob("rpt_public_*.sql")
    }
    assert not any(name.startswith("rpt_public") for name in EXPORT.LOCAL_TABLES)


@needs_warehouse
def test_full_export_on_the_real_warehouse(tmp_path):
    assert EXPORT.main(["--out", str(tmp_path), "--no-excel"]) == 0
    EXPORT.assert_public_only(tmp_path / "public")
    local = {
        path.relative_to(tmp_path / "local").as_posix()
        for path in (tmp_path / "local").rglob("*.parquet")
    }
    assert {f"{name}.parquet" for name in EXPORT.LOCAL_TABLES} <= local
    assert {f"analysis/{result}.parquet" for result in EXPORT.ANALYSIS_IN_MODEL} <= local
    company = pd.read_parquet(tmp_path / "local" / "dim_company.parquet")
    assert company["in_ghg_balanced_panel"].sum() == 38  # the balanced panel of Q1 and Q5


# ------------------------------------------------------------------ Power BI model


def _tmdl_files(project: str) -> dict[str, str]:
    folder = POWERBI / project / f"{project}.SemanticModel" / "definition"
    return {
        path.relative_to(folder).as_posix(): path.read_text(encoding="utf-8")
        for path in folder.rglob("*.tmdl")
    }


def _measures(project: str) -> dict[str, str]:
    """measure name -> DAX expression, read back from the committed TMDL."""
    text = _tmdl_files(project)["tables/_Measures.tmdl"]
    measures: dict[str, list[str]] = {}
    name = None
    for line in text.splitlines():
        match = re.match(r"\tmeasure ('(?:[^']|'')+'|\w+) =\s*$", line)
        if match:
            name = match.group(1).strip("'").replace("''", "'")
            measures[name] = []
        elif name and line.startswith("\t\t\t"):
            measures[name].append(line.strip())
        else:
            name = None
    return {key: "\n".join(lines) for key, lines in measures.items()}


@pytest.mark.parametrize("project", ["esg_lens_local", "esg_lens_public"])
def test_power_bi_project_is_complete(project):
    base = POWERBI / project
    pbip = json.loads((base / f"{project}.pbip").read_text(encoding="utf-8"))
    assert pbip["artifacts"] == [{"report": {"path": f"{project}.Report"}}]
    pbir = json.loads((base / f"{project}.Report" / "definition.pbir").read_text(encoding="utf-8"))
    assert pbir["datasetReference"]["byPath"]["path"] == f"../{project}.SemanticModel"
    files = _tmdl_files(project)
    assert {"database.tmdl", "model.tmdl", "expressions.tmdl", "relationships.tmdl"} <= set(files)
    assert "discourageImplicitMeasures" in files["model.tmdl"]
    assert "crossFilteringBehavior" not in files["relationships.tmdl"]  # single direction
    assert "sortByColumn:" in files["tables/dim_fiscal_year.tmdl"]
    measures = files["tables/_Measures.tmdl"]
    assert measures.count("\tmeasure ") == measures.count("displayFolder:") > 20
    for table, text in files.items():
        if table.startswith("tables/") and table != "tables/_Measures.tmdl":
            assert "\tmeasure " not in text, f"{table}: measures belong in _Measures"


def test_no_local_path_is_committed_in_the_power_bi_projects():
    drive = re.compile(r"\b[A-Za-z]:\\")
    for path in POWERBI.rglob("*"):
        if path.is_file():
            text = path.read_text(encoding="utf-8", errors="replace")
            assert not drive.search(text), (
                f"{path}: hard-coded path (use the ExportFolder parameter)"
            )
    for project in ("esg_lens_local", "esg_lens_public"):
        expressions = _tmdl_files(project)["expressions.tmdl"]
        assert 'expression ExportFolder = "SET-ME' in expressions
        assert "IsParameterQuery = true" in expressions


def test_no_measure_averages_an_intensity():
    average = re.compile(r"AVERAGEX?\s*\(\s*[^)]*(intensity|_pct|share|_per_)", re.IGNORECASE)
    for project in ("esg_lens_local", "esg_lens_public"):
        for name, dax in _measures(project).items():
            assert not average.search(dax), f"{project} {name}: average of an intensity or share"


def test_local_ratio_measures_follow_the_catalogue():
    """Every ratio of sums of the semantic layer has a measure that sums the same numerator
    and denominator over the rows that have both."""
    with (ROOT / "dbt" / "seeds" / "kpi_catalogue.csv").open(encoding="utf-8") as fh:
        ratios = [row for row in csv.DictReader(fh) if row["metric_type"] == "ratio_of_sums"]
    measures = _measures("esg_lens_local")
    for row in ratios:
        matching = [
            dax
            for dax in measures.values()
            if f"SUMX ( rows_with_both, fct_company_year[{row['numerator']}] )" in dax
            and f"SUMX ( rows_with_both, fct_company_year[{row['denominator']}] )" in dax
            and f"fct_company_year[{row['denominator']}] <> 0" in dax
        ]
        assert matching, f"no measure for {row['kpi_name']}"
    assert len(ratios) == len([d for d in measures.values() if "rows_with_both" in d])


def test_public_measures_never_recompute_a_gated_statistic():
    """The public model reads group figures from rpt_public_group_metrics; from company rows it
    only picks one value (SELECTEDVALUE) or counts companies."""
    forbidden = re.compile(
        r"\b(SUM|SUMX|AVERAGE|AVERAGEX|MEDIAN|MEDIANX|DIVIDE|MIN|MINX|PERCENTILE\w*)\s*\("
    )
    for name, dax in _measures("esg_lens_public").items():
        assert not forbidden.search(dax), f"{name}: computes a statistic"
        for table in re.findall(r"MAX \( (\w+)\[", dax):
            assert table == "rpt_public_group_metrics", f"{name}: MAX over {table}"
        tables = set(re.findall(r"\b(\w+)\[", dax))
        assert tables <= {
            "rpt_public_group_metrics",
            "rpt_public_company_year",
            "rpt_public_disclosure_quality",
            "rpt_public_concentration",
            "rpt_public_company_percentile",
            "dim_fiscal_year",
            "dim_company",
        }, name


def test_public_model_reads_only_public_files():
    for table, text in _tmdl_files("esg_lens_public").items():
        for file in re.findall(r'ExportFolder & "\\\\([^"]+)"', text):
            assert file in EXPORT.PUBLIC_FILES, f"{table} reads {file}"


@needs_warehouse
def test_committed_power_bi_files_are_generated_from_the_warehouse():
    assert PBI.main(["--check"]) == 0


@pytest.mark.skipif(
    sys.platform != "win32" or not shutil.which("powershell"), reason="Windows only"
)
def test_tmdl_parses_with_power_bi_desktop():
    result = subprocess.run(
        [
            "powershell",
            "-NoProfile",
            "-ExecutionPolicy",
            "Bypass",
            "-File",
            str(SCRIPTS / "check_tmdl.ps1"),
        ],
        capture_output=True,
        text=True,
        check=False,
    )
    if "library not found" in result.stderr:
        pytest.skip("Power BI Desktop is not installed")
    assert result.returncode == 0, result.stdout + result.stderr
    models = {entry["model"]: entry for entry in json.loads(result.stdout)}
    for entry in models.values():
        assert (
            entry["ok"]
            and entry["measures_without_description"] == []
            and entry["bidirectional"] == []
        )


# ------------------------------------------------------------------ Excel


def _public_workbook(tmp_path: Path) -> Path:
    public = EXPORTS / "public"
    if not (public / "rpt_public_group_metrics.parquet").exists():
        pytest.skip("run dev.ps1 bi-export first")
    return EXCEL.build(public, tmp_path / "esg_lens_public.xlsx")


def test_workbook_structure_and_formulas(tmp_path):
    workbook = load_workbook(_public_workbook(tmp_path))
    assert workbook.sheetnames[:5] == [
        "README",
        "Data",
        "Peer Benchmark",
        "Sector Summary",
        "DQS view",
    ]
    assert "tblCompanyYear" in workbook["Data"].tables
    peer = workbook["Peer Benchmark"]
    assert len(peer.data_validations.dataValidation) == 3
    for cell in ("B9", "B10", "B12", "B13", "B14", "B15", "B16", "B17", "B18"):
        assert str(peer[cell].value).startswith("="), f"{cell} must be a formula"
    assert "_xlfn.XLOOKUP(" in peer["B10"].value and "tblPercentile" in peer["B15"].value
    assert "COUNTIFS(" in peer["F15"].value and "F15" in peer["B19"].value
    assert "tblPercentile" in workbook["Percentile Data"].tables
    summary = workbook["Sector Summary"]
    values = [
        c.value
        for row in summary.iter_rows(min_row=6, min_col=3)
        for c in row
        if c.value is not None
    ]
    assert values and all(str(v).startswith("=") for v in values), (
        "Sector Summary must hold formulas only"
    )
    assert any('"n<5"' in str(v) for v in values)
    readme = " ".join(str(c.value) for c in workbook["README"]["A"] if c.value)
    assert "not an ESG rating" in readme and "not investment advice" in readme
    assert "http" not in readme.lower()


def _excel_check(workbook: Path, cases: list[dict], tmp_path: Path) -> dict:
    if sys.platform != "win32":
        pytest.skip("Excel check runs on Windows")
    probe = subprocess.run(
        [
            "powershell",
            "-NoProfile",
            "-Command",
            "try { $x = New-Object -ComObject Excel.Application; $x.Quit(); 'yes' } catch { 'no' }",
        ],
        capture_output=True,
        text=True,
        check=False,
    )
    if probe.stdout.strip() != "yes":
        pytest.skip("Excel is not installed")
    (tmp_path / "cases.json").write_text(json.dumps(cases), encoding="utf-8")
    out = tmp_path / "excel.json"
    result = subprocess.run(
        ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(SCRIPTS / "check_excel.ps1"),
         "-Path", str(workbook), "-CasesJson", str(tmp_path / "cases.json"), "-Out", str(out)],
        capture_output=True,
        text=True,
        check=False,
    )  # fmt: skip
    assert result.returncode == 0, result.stdout + result.stderr
    return json.loads(out.read_text(encoding="utf-8-sig"))


@needs_warehouse
def test_workbook_calculates_in_excel_and_agrees_with_the_warehouse(tmp_path):
    """The workbook is built from exports/public, so warehouse figures are compared rounded to
    the decimals of the public export."""

    def public(value: float) -> float:
        return round(value, EXPORT.PUBLIC_DECIMALS)

    workbook = _public_workbook(tmp_path)
    con = duckdb.connect(str(WAREHOUSE), read_only=True)
    try:
        expected = con.sql(
            """
            select company.display_name, percentile.kpi_name, percentile.fiscal_year_label,
                   percentile.percentile_in_group, percentile.ranking_group,
                   percentile.percentile_null_reason as reason_not_ranked
            from fct_company_kpi_percentile as percentile
            inner join dim_company as company on percentile.isin = company.isin
            where percentile.kpi_name in ('ghg_intensity_tco2e_per_cr', 'water_intensity_kl_per_cr')
              and percentile.fiscal_year_label in ('FY2024-25', 'FY2025-26')
              and percentile.value is not null
            """
        ).df()
        groups = con.sql("select * from rpt_public_group_metrics").df()
    finally:
        con.close()
    labels = {column: label for label, column, *_ in EXCEL.BENCHMARK_KPIS}
    cases = [
        {"company": row.display_name, "kpi": labels[row.kpi_name], "year": row.fiscal_year_label}
        for row in expected.itertuples()
    ]
    result = _excel_check(workbook, cases, tmp_path)

    errors = {sheet["name"]: sheet["errors"] for sheet in result["sheets"]}
    assert set(errors.values()) == {0}, errors
    assert sum(sheet["formulas"] for sheet in result["sheets"]) > 900

    medians = groups[groups["group_type"] == "ranking_group"].set_index(
        ["group_value", "fiscal_year_label"]
    )
    by_case = {(c["company"], c["kpi"], c["year"]): c for c in result["peer_cases"]}
    ranked = not_material = 0
    for row in expected.itertuples():
        case = by_case[(row.display_name, labels[row.kpi_name], row.fiscal_year_label)]
        assert case["ranking_group"] == row.ranking_group
        if pd.isna(row.percentile_in_group):
            assert case["ranked"] != "yes" and case["check"] == "-", case
        else:
            assert case["ranked"] == "yes", case
            assert case["percentile"] == pytest.approx(public(row.percentile_in_group), abs=1e-9), (
                case
            )
            # the hidden COUNTIFS formulas give the same rank and percentile
            assert case["check"].startswith("OK"), case
            ranked += 1
        if row.reason_not_ranked == "not_material":
            assert case["ranked"] == "not ranked for this group (not material)", case
            not_material += 1
        median = medians.loc[
            (row.ranking_group, row.fiscal_year_label),
            row.kpi_name.replace("_tco2e_per_cr", "_median").replace("_kl_per_cr", "_median"),
        ]
        if pd.isna(median):
            assert case["median"] == "n<5", case
        else:
            assert case["median"] == pytest.approx(public(median), rel=1e-12), case
    assert ranked > 40 and not_material > 0

    columns = {header: column for header, column, *_ in EXCEL.SUMMARY_COLUMNS}
    table = groups.set_index(["group_type", "group_value", "fiscal_year_label"])
    shown = withheld = 0
    for cell in result["summary"]:
        key = (cell["group_type"], cell["group_value"], cell["year"])
        if key not in table.index:
            assert cell["value"] == "-", cell
            continue
        value = table.loc[
            (cell["group_type"], cell["group_value"], cell["year"]), columns[cell["header"]]
        ]
        if pd.isna(value):
            assert cell["value"] in ("n<5", "not yet observable"), cell
            withheld += 1
        else:
            assert cell["value"] == pytest.approx(public(float(value)), rel=1e-12), cell
            shown += 1
    assert shown > 500 and withheld > 0


# ------------------------------------------------------------------ reconciliation


def _semantic_ready() -> bool:
    return (
        _real_warehouse()
        and (ROOT / "dbt" / "target" / "semantic_manifest.json").exists()
        and (EXPORTS / "local" / "fct_company_year.parquet").exists()
    )


@pytest.mark.skipif(
    not _semantic_ready(), reason="needs the warehouse, dbt parse and dev.ps1 bi-export"
)
def test_dax_formulas_reconcile_with_mf_query(tmp_path):
    tables = RECONCILE.load_local()
    ours = RECONCILE.dax_equivalent(tables)
    compared, largest, problems = RECONCILE.compare(ours, RECONCILE.mf_values(tmp_path))
    assert not problems, problems[:10]
    assert compared >= 80 and largest < 1e-9
    public_compared, public_problems = RECONCILE.public_matches(ours)
    assert not public_problems and public_compared > 50
    assert len(RECONCILE.CHECKS) == 6


@pytest.mark.skipif(
    not (EXPORTS / "local" / "fct_company_year.parquet").exists(), reason="run dev.ps1 bi-export"
)
def test_checklist_in_the_docs_holds_the_current_values():
    text = RECONCILIATION_DOC.read_text(encoding="utf-8")
    for item in RECONCILE.checklist(RECONCILE.load_local()):
        expected = f"{item['expected']:{'.0f' if item['format'] == '0' else ('.1f' if item['format'] == '0.0' else '.2f')}}"
        assert f"| {item['measure']} |" in text, item["measure"]
        assert f"**{expected}**" in text, f"{item['measure']}: expected {expected} in the checklist"
    assert text.count("- [ ]") >= 6


def test_build_guide_covers_five_pages():
    text = GUIDE.read_text(encoding="utf-8")
    headings = [line for line in text.splitlines() if line.startswith("## Page")]
    assert len(headings) == 5
    for page in (
        "Overview",
        "Decarbonisation",
        "Concentration",
        "Disclosure quality",
        "Company card",
    ):
        assert any(page in heading for heading in headings), page
    assert text.count("**Interview talking point.**") == 5
    for word in ("colour", "provisional", "screenshot"):
        assert word in text.lower()
    measures = set(_measures("esg_lens_local")) | set(_measures("esg_lens_public"))
    for name in re.findall(r"`\[([^\]]+)\]`", text):
        assert name in measures, f"the guide names a measure that does not exist: {name}"


def test_reconciliation_values_are_derived():
    """The checklist is public: intensities, shares, medians and counts only."""
    text = RECONCILIATION_DOC.read_text(encoding="utf-8")
    assert not re.search(r"\d[\d,.]*\s*(tCO2e|GJ|kl)\b(?!\s*per)", text)


@needs_warehouse
def test_public_concentration_matches_week_10_q2():
    """rpt_public_concentration gives the Q2 shares: top 1, top 5, top 10 and companies for 80%."""
    sys.path.insert(0, str(ROOT / "tests"))
    from test_analysis import QUERIES, RUNNER

    con = duckdb.connect(str(WAREHOUSE), read_only=True)
    try:
        q2 = RUNNER.run_queries(
            con,
            [
                q
                for q in QUERIES
                if q.result_id == "q02_emission_concentration__concentration_by_year"
            ],
        )
        concentration = con.sql("select * from rpt_public_concentration").df()
    finally:
        con.close()
    q2 = q2["q02_emission_concentration__concentration_by_year"].set_index("fiscal_year_label")
    assert len(q2) == 4
    for year, rows in concentration.groupby("fiscal_year_label"):
        by_rank = rows.set_index("emitter_rank")["cumulative_share_pct"]
        assert round(by_rank[1], 1) == q2.loc[year, "top_1_share_pct"]
        assert round(by_rank[5], 1) == q2.loc[year, "top_5_share_pct"]
        assert round(by_rank[10], 1) == q2.loc[year, "top_10_share_pct"]
        assert (
            int((rows["cumulative_share_pct"] < 80).sum()) + 1
            == q2.loc[year, "companies_for_80_pct"]
        )
        assert len(rows) == q2.loc[year, "n_companies"]
    assert q2["top_5_share_pct"].between(85.5, 86.2).all()
    assert not {c for c in concentration.columns if c.endswith(("_tco2e", "total"))}


@pytest.mark.skipif(
    not (EXPORTS / "public" / "rpt_public_concentration.parquet").exists(),
    reason="run dev.ps1 bi-export",
)
def test_public_checklist_in_the_guide_holds_the_current_values():
    text = GUIDE.read_text(encoding="utf-8")
    for item in RECONCILE.public_checklist():
        if item["format"] == "text":
            expected = str(item["expected"])
        else:
            digits = {"0": ".0f", "0.0": ".1f", "0.00": ".2f"}[item["format"]]
            expected = f"{float(item['expected']):{digits}}"
        assert f"| {item['measure']} |" in text, item["measure"]
        assert f"**{expected}**" in text, f"{item['measure']}: expected {expected} in the guide"
